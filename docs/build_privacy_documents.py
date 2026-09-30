from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
LOGO = ROOT.parent / "mobile/interpath_mobile/assets/images/interpathmed_logo.png"

DOCUMENTS = [
    (
        ROOT / "privacy-policy.md",
        ROOT / "Interpath_Results_Privacy_Policy.docx",
        "Privacy Policy",
    ),
    (
        ROOT / "data-deletion.md",
        ROOT / "Interpath_Results_Data_Deletion_Request.docx",
        "Data Deletion Request",
    ),
]


def set_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def shade_run(run, fill="FFF2CC"):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    run._element.get_or_add_rPr().append(shading)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)
    set_font(run, size=9, color=(96, 96, 96))


def configure_document(doc, footer_label):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(35, 43, 51)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_before = Pt(6)
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.keep_with_next = True
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 15, 14, 6),
        ("Heading 2", 11.5, 9, 4),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    for list_style in ("List Bullet", "List Number"):
        style = styles[list_style]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.27)
        style.paragraph_format.first_line_indent = Inches(-0.18)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.08

    footer = section.footer
    footer.distance = Inches(0.3)
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"Interpath Results  |  {footer_label}  |  ")
    set_font(run, size=9, color=(96, 96, 96))
    add_page_field(paragraph)


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\](?:\([^\)]*\))?)")


def add_inline(paragraph, text):
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            add_segment(paragraph, text[position:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, bold=True)
        elif token.startswith("["):
            label_end = token.find("]")
            label = token[1:label_end]
            suffix = token[label_end + 1:]
            if suffix.startswith("(") and suffix.endswith(")"):
                target = suffix[1:-1]
                rendered = label if not target else f"{label} ({target})"
            else:
                rendered = token
            add_segment(paragraph, rendered)
        position = match.end()
    if position < len(text):
        add_segment(paragraph, text[position:])


def add_segment(paragraph, text):
    if not text:
        return
    parts = re.split(r"(\[[^\]]*(?:insert|INSERT)[^\]]*\])", text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run(part)
        set_font(run)
        if part.startswith("[") and re.search(r"insert", part, re.IGNORECASE):
            set_font(run, bold=True, color=(127, 84, 0))
            shade_run(run)


def add_logo(doc):
    if not LOGO.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(5)
    picture = paragraph.add_run().add_picture(str(LOGO), width=Inches(1.6))
    picture._inline.docPr.set("name", "Interpath Medical Laboratories logo")
    picture._inline.docPr.set("descr", "Interpath Medical Laboratories logo")


def build_document(source, destination, footer_label):
    doc = Document()
    configure_document(doc, footer_label)
    add_logo(doc)

    lines = source.read_text(encoding="utf-8").splitlines()
    first_title = True
    for raw in lines:
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, line[2:])
            first_title = False
        elif line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            add_inline(paragraph, line[3:])
        elif line.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            add_inline(paragraph, line[4:])
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
        elif re.match(r"^\d+\. ", line):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.27)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.line_spacing = 1.08
            add_inline(paragraph, line)
        else:
            paragraph = doc.add_paragraph()
            if line.startswith("**Effective date:**") or line.startswith("**Last updated:**"):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(7)
                paragraph.paragraph_format.keep_with_next = True
            add_inline(paragraph, line.replace("  ", " "))

    core = doc.core_properties
    core.title = lines[0].removeprefix("# ") if lines else footer_label
    core.subject = "Interpath Results privacy and data protection"
    core.author = "Interpath Medical Laboratories"
    core.keywords = "privacy, data protection, laboratory results, WhatsApp"
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


for source, destination, footer_label in DOCUMENTS:
    build_document(source, destination, footer_label)
    print(destination)
